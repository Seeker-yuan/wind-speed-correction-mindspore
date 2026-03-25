from docx import Document

src = r"c:\Users\31876\Desktop\一种基于多视图拓扑与掩码图神经网络的风机风速缺失数据同步推演方法.docx"
out = "patent_docx_tables.txt"

doc = Document(src)
lines = []
for ti, table in enumerate(doc.tables, start=1):
    lines.append(f"[Table {ti}]")
    for ri, row in enumerate(table.rows, start=1):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        lines.append(f"R{ri:02d} | " + " | ".join(cells))
    lines.append("")

with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"tables={len(doc.tables)} out={out}")
